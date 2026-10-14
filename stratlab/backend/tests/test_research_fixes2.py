"""Company research fixes: negative net worth, ADRs reporting in rupees, capex years, the deck's full checklist,
dividend yield, the share card's stale figures and the report card's met rule."""
import io
from datetime import date

from pptx import Presentation

from app import checklist, company_cards, deck, deepdive, industry, report_card as rc
from app.intel import sec
from app.intel.screener import summary
from tests.test_deepdive import company
from tests.test_investor import view_for_deck


def test_negative_net_worth_has_no_debt_to_equity_or_pb():
    p = company()
    p["balance"]["rows"].update({"Reserves": [-500, -600], "Equity Capital": [100, 100], "Borrowings": [900, 1000]})
    p["balance"]["cols"] = p["balance"]["cols"][-2:]
    p["ratios"] = {"Current Price": "100", "Book Value": "-80"}
    s = summary(p)
    assert s["debt_equity"] is None and s["pb"] is None and s["net_worth_negative"]
    row = {c["label"]: c for c in checklist.evaluate(p, deepdive.numbers(p))["checks"]}["Debt to equity"]
    assert row["state"] == "watch" and "net worth is negative" in row["value"]


def adr():
    return {"region": "US", "currency": "INR", "unit": "INR million", "fx": {"rate": 0.012},
            "ratios": {"Market Cap": "10000", "Stock P/E": "20"},
            "pl": {"cols": ["Mar 2025", "TTM"], "rows": {"Net Profit": [48000, 50000], "Operating Profit": [60000, 62000]}}}


def test_an_inr_reporting_adr_is_converted_like_any_foreign_filer():
    note = deepdive.profit_note(adr())          # 50,000 m INR is $600 m against $500 m of earnings: 1.2x, no note
    assert note is None
    p = adr()
    p["pl"]["rows"]["Net Profit"][-1] = 100000   # $1,200 m: 2.4x the $500 m the P/E rests on
    assert "2.4×" in deepdive.profit_note(p)
    v = industry.valuation({**adr(), "balance": {"cols": ["Mar 2025"], "rows": {"Borrowings": [0]}}},
                           {"market_cap_cr": 10000, "pe": None}, "general", "ev_ebitda")
    assert v["pe"] == 10000 / 0.012 / 50000      # market value in INR over INR profit


def test_three_year_capex_says_how_many_years_were_summed():
    n = deepdive.numbers(company())
    assert n["capex_years"] == 3 and n["capex_3y_total"] == 62
    p = company()
    p["balance"]["rows"]["Fixed Assets"][-2] = 40          # fixed assets fell in FY24: that year's capex is cut
    n = deepdive.numbers(p)
    assert n["capex_years"] == 2


def test_the_deck_shows_every_checklist_row():
    v = view_for_deck()
    row = v["checklist"]["checks"][0]
    v["checklist"]["checks"] = [dict(row, label=f"Check {i}") for i in range(16)]
    prs = Presentation(io.BytesIO(deck.build(v)))
    cells = [c.text for s in prs.slides for sh in s.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
    assert all(f"Check {i}" in cells for i in range(16))


def test_us_dividend_yield_ignores_an_older_year_when_the_latest_is_blank():
    p = {"region": "US", "currency": "USD", "ratios": {}, "shares": 1e8,
         "cashflow": {"cols": ["Dec 2023", "Dec 2024"], "rows": {"Dividends paid": [-500, None]}},
         "balance": {"cols": ["Dec 2024"], "rows": {"Equity": [1000]}}, "pl": {"cols": ["Dec 2024"], "rows": {}}}
    assert sec.ratios(p, 10.0)["Dividend Yield"] == 0.0        # not the 500 of 2023


def test_the_special_dividend_note_uses_the_companys_currency():
    from app import stock_pages
    f = {"currency": "USD", "region": "US", "div_special": 1.5, "div_yield": 3.0, "price": 100}
    text = " ".join(v for _, v in stock_pages._stats(f))
    assert "special dividend of $1.50" in text


def test_the_share_card_drops_figures_the_page_marks_stale():
    f = {"roe": 20.0, "net_margin": 10.0, "opm": 15.0, "price_at": "2026-09-01", "years": [{"year": "Mar 2025"}],
         "growth": {"sales_cagr_3y": 12.0}, "region": "IN"}
    labels = [k for k, _ in company_cards.card(f, "IN", "X")["facts"]]
    assert "Return on equity" not in labels and "Net margin" not in labels
    fresh = dict(f, price_at="2025-06-01")
    assert "Return on equity" in [k for k, _ in company_cards.card(fresh, "IN", "X")["facts"]]


def test_met_means_at_or_above_the_target_less_the_tolerance():
    nums = deepdive.numbers(company())       # FY25 sales growth 17.6%
    def result(low, high):
        g = {"metric": "revenue_growth", "low": low, "high": high, "period": "FY25", "what": "x", "quote": "", "said_at": "2024-05-10T10:00",
             "source": {"title": "t", "at": "2024-05-10T10:00", "url": "u"}}
        return rc.view({"guidance": [g], "at": "x"}, nums, today=date(2026, 9, 30))["rows"][0]["result"]
    assert result(17.0, None) == "met" and result(10.0, None) == "met"          # beating the target is met
    assert result(19.0, None) == "missed" and result(15.0, 18.0) == "met" and result(18.0, 20.0) == "met"
