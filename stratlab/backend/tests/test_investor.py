"""Phase 6: the investor checklist's fixed rules, the company deck (a real PowerPoint file) and the investor home."""
import io
import json

from pptx import Presentation

from app import checklist, db, deck, deepdive, investor, report_card
from tests.test_deepdive import api, company  # noqa: F401  (api is a fixture)


def with_holding(p, promoters):
    p["shareholding"] = {"cols": [f"Q{i}" for i in range(len(promoters))], "rows": {"Promoters": promoters, "FIIs": [10] * len(promoters)}}
    p["ratios"] = {"ROCE": "18.2 %", "Stock P/E": "24", "Market Cap": "₹ 1,200 Cr."}
    return p


def states(res):
    return {c["label"]: (c["state"], c["value"]) for c in res["checks"]}


def test_checklist_rules():
    p = with_holding(company(), [55, 55, 54, 53, 52.5])
    nums = deepdive.numbers(p)
    trend = {"stage": 2, "st_up": True, "st_days": 12}
    fsum = {"red": 1, "amber": 0, "fund_raise": True}
    card = {"met": 1, "missed": 3, "score": 25}
    s = states(checklist.evaluate(p, nums, fsum, trend, card))
    assert s["Price in Stage 2"][0] == "pass" and s["Supertrend up"] == ("pass", "Up for 12 days")
    assert s["Sales growth, 3 years"][0] == "pass"            # 130 → 200 over 3 years: ~15% a year
    assert s["Return on capital employed"] == ("pass", "18.2%")
    assert s["Operating margin holding up"][0] == "pass"      # 19 now vs 17 three years before
    assert s["Debt to equity"][0] == "na"                     # no reserves row: can't compute
    assert s["Profit turning into cash"][0] == "pass"         # (20+24+30) / (18+20+25) = 1.17
    assert s["Promoter holding"] == ("pass", "52.5%") and s["Promoter holding over a year"][0] == "fail"   # -2.5 points
    assert s["Red-flag filings, last 3 months"][0] == "fail" and s["Fund raise filed, last 3 months"][0] == "watch"
    assert s["Delivered on targets"] == ("fail", "1 of 4 met")


def test_checklist_without_optional_inputs():
    res = checklist.evaluate({"pl": None}, deepdive.numbers({}))
    labels = {c["label"] for c in res["checks"]}
    assert "Price in Stage 2" not in labels and "Promoter holding" not in labels and "Delivered on targets" not in labels
    assert all(c["state"] == "na" for c in res["checks"] if c["label"] != "Free cash flow, 3 years") and res["scored"] == 0


def view_for_deck():
    p = with_holding(company(), [55, 55, 54, 53, 52.5])
    nums = deepdive.numbers(p)
    card = report_card.view({"guidance": [{"metric": "margin", "low": 18.0, "high": None, "period": "FY25", "what": "18% margin",
                                           "quote": "about 18%", "said_at": "2024-05-10T10:00",
                                           "source": {"title": "Call", "at": "2024-05-10T10:00", "url": "https://x/1.pdf"}}], "at": "x"}, nums)
    return {"symbol": "ACME", "name": "Acme Industries", "about": p["about"], "numbers": nums, "snapshot": {"market_cap_cr": 1200, "roce": 18.2},
            "documents": [{"kind": "transcript", "at": "2026-08-10T10:00", "title": "Call", "url": "https://x/1.pdf"}],
            "reads": {"business": {"summary": "Pumps.", "segments": [{"name": "Water", "share_pct": 60, "what": ""}], "customers": "Utilities",
                                   "drivers": ["Tenders"], "strengths": [], "risks": ["Steel"], "sources": []},
                      "plans": {"capex": [{"what": "New plant", "amount": "Rs 500 crore", "timeline": "FY28", "status": "planned", "quote": "",
                                           "source": {"title": "Deck", "at": "2026-08-01T10:00", "url": "https://x/2.pdf", "kind": "presentation"}}],
                                "outlook": [{"statement": "15% growth", "quote": "we expect 15%", "source": None}], "sources": []}},
            "card": card, "checklist": checklist.evaluate(p, nums, {"red": 0, "fund_raise": False}, None, card)}


def test_deck_is_a_real_presentation():
    prs = Presentation(io.BytesIO(deck.build(view_for_deck())))
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text), "") for s in prs.slides]
    assert titles == ["Acme Industries", "Sales and net profit", "The last eight quarters", "Capex and cash", "Business model",
                      "Capex and growth plans, in management's words", "Management report card", "Investor checklist", "Sources"]
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Not investment advice" in text and "Screener" not in text
    assert deck._cr(1234567) == "12,34,567" and deck._cr(-950) == "-950" and deck._cr(None) == "–"


def test_deck_without_reads_still_builds():
    v = view_for_deck()
    v.update(reads=None, card=None, checklist=None)
    prs = Presentation(io.BytesIO(deck.build(v)))
    assert len(prs.slides) == 6
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Not in this deck yet" in text and "Read the documents" in text and "Check past calls" in text


def test_sector_of_prefers_main_sectors():
    assert investor.sector_of("IN", "HDFCBANK") == "NIFTY BANK"
    assert investor.sector_of("IN", "NOT-A-STOCK") is None


def test_deep_view_has_checklist_and_deck_and_investor_home(api):  # noqa: F811
    c, who, _calls, _usage = api
    assert c.get("/research/investor").status_code == 402 and c.get("/research/deep/ACME/deck").status_code == 402
    who["p"] = {"id": "u1", "plan": "pro", "_plan": "pro"}
    v = c.get("/research/deep/ACME").json()
    assert v["checklist"]["counts"]["pass"] >= 3 and v["filings"]["red"] == 0 and "snapshot" in v
    r = c.get("/research/deep/ACME/deck")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('ACME-deep-dive.pptx"')
    assert len(Presentation(io.BytesIO(r.content)).slides) >= 5
    assert c.get("/research/investor").json()["rows"] == []
    db.set_setting("watchlist:u1", json.dumps({"items": [{"region": "IN", "symbol": "ACME"}, {"region": "IN", "symbol": "NOPE"},
                                                          {"region": "US", "symbol": "AAPL"}]}))
    rows = c.get("/research/investor").json()["rows"]
    assert [x["symbol"] for x in rows] == ["ACME", "NOPE"]
    assert rows[0]["name"] == "Acme Industries" and rows[0]["checks"]["pass"] >= 3 and rows[0]["red"] == 0
    assert rows[1]["checks"] is None and "Screener" not in (rows[1]["problem"] or "")


def test_bank_deck_skips_capex_and_uses_lender_labels():
    v = view_for_deck()
    v["numbers"] = {**v["numbers"], "bank": True}
    v["snapshot"] = {"roe": 16.5, "div_yield": 1.2}
    prs = Presentation(io.BytesIO(deck.build(v)))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Capex and cash" not in text and "Revenue and net profit" in text and "ROE" in text and "Debt / equity" not in text
