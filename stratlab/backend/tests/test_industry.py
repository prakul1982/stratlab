"""Industry-aware checks: which rule set and which operating measures fit a company, the classification read from the
company page, the measures pulled from its presentation, and the deck slide that shows them."""
import io
import json

from pptx import Presentation

from app import checklist, deck, deepdive, industry
from app.intel.screener import parse
from tests.test_deepdive import FakeDocs, company
from tests.test_investor import view_for_deck


def labels(p, sym=None):
    res = checklist.evaluate(p, deepdive.numbers(p), symbol=sym)
    return res["industry"]["group"], {c["label"]: c for c in res["checks"]}


def test_classification_from_page_name_statements_and_lists():
    assert industry.classify({"industry_path": ["Financial Services", "Insurance", "Life Insurance"]})["group"] == "insurer"
    assert industry.classify({"name": "Bajaj Holdings & Investment Ltd"})["group"] == "holding"
    assert industry.classify({"industry_path": ["Utilities", "Power", "Power Generation"]})["group"] == "utility"
    assert industry.classify({"industry_path": ["Realty", "Residential, Commercial Projects"]})["group"] == "realty"
    assert industry.classify({"industry_path": ["Commodities", "Cement & Cement Products"]})["group"] == "cyclical"
    assert industry.classify({"industry_path": ["Capital Goods", "Electrical Equipment"]})["group"] == "general"
    assert industry.classify({"name": "Some Co"}, {"bank": True})["group"] == "lender"     # the statements say lender
    assert industry.classify({"name": "X"}, None, "DLF")["group"] == "realty"              # StratLab's sector lists
    assert industry.classify({"name": "X"}, None, "NTPC")["group"] == "utility"
    assert industry.classify({"name": "X"}, None, "TCS")["group"] == "general"


def test_page_classification_is_parsed():
    html = """<html><h1>Apollo Hospitals Enterprise Ltd</h1><section id="peers"><p class="sub">
      <a href="/market/IN05/" title="Broad Sector">Healthcare</a> <a href="/market/IN05/IN0501/" title="Sector">Healthcare</a>
      <a href="/market/IN05/IN0501/IN050101/" title="Industry">Healthcare Services</a>
      <a href="/market/IN05/IN0501/IN050101/IN050101001/" title="Basic Industry">Hospital</a>
      <a href="/company/compare/">Edit columns</a></p></section></html>"""
    p = parse(html)
    assert p["industry_path"] == ["Healthcare", "Healthcare", "Healthcare Services", "Hospital"]
    assert industry.measures(p)["key"] == "hospital"
    assert any("ARPOB" in m for m in industry.measures(p)["measures"])


def test_rules_soften_for_the_industry_that_would_always_fail():
    p = company()
    p["balance"]["rows"]["Reserves"] = [5] * 6           # tiny equity: debt to equity far above 1
    p["balance"]["rows"]["Equity Capital"] = [1] * 6
    p["industry_path"] = ["Utilities", "Power", "Power Generation"]
    grp, c = labels(p)
    assert grp == "utility" and c["Debt to equity"]["state"] == "watch" and "Shown as watch for power" in c["Debt to equity"]["rule"]
    p["industry_path"] = ["Capital Goods", "Industrial Machinery"]
    grp, c = labels(p)
    assert grp == "general" and c["Debt to equity"]["state"] == "fail"


def test_insurers_and_holding_companies_use_return_on_equity():
    p = company()
    p["ratios"] = {"ROE": "12 %", "ROCE": "30 %"}
    p["industry_path"] = ["Financial Services", "Insurance"]
    grp, c = labels(p)
    assert grp == "insurer" and c["Return on equity"]["state"] == "watch"   # 12%: between 8 and 14
    assert not {"Return on capital employed", "Debt to equity", "Operating margin holding up"} & c.keys()
    p.pop("industry_path")
    p["name"] = "Tata Investment Corporation Ltd"
    grp, c = labels(p)
    assert grp == "holding" and c["Return on equity"]["state"] == "pass"   # holding: pass at 10%


def test_measures_by_industry():
    assert industry.measures({"industry_path": ["Consumer Services", "Hotels & Resorts"]})["label"] == "Hotels"
    assert industry.measures({"name": "X"}, "INFY")["key"] == "it"
    assert industry.measures({"name": "X"}, "HDFCBANK")["measures"][0].startswith("Net interest margin")
    assert industry.measures({"name": "Unknown Widgets"})["key"] == "general"


def test_presentation_read_pulls_the_measures(monkeypatch):
    seen = []

    def complete(system, text, **kw):
        seen.append(text)
        if "makes money" in system:
            return json.dumps({"summary": "Hospitals.", "segments": [], "measures": [
                {"name": "ARPOB", "value": "Rs 62,000 a day", "period": "Q1 FY27", "change": "+8% YoY", "quote": "ARPOB rose to 62,000", "source": "S1"},
                {"name": "Occupancy", "value": "", "source": "S1"},                       # no value: dropped
                {"name": "Beds", "value": "10,000", "source": "S7"}]})                    # no quote in any document: dropped
        return json.dumps({"capex": [], "outlook": []})
    monkeypatch.setattr(deepdive, "complete", complete)
    docs = [{"kind": "presentation", "at": "2026-08-01T10:00", "title": "Investor Presentation", "url": "u1"}]
    out = deepdive.read("APOLLOHOSP", "Apollo", "Hospitals", docs, FakeDocs({"u1": "ARPOB rose to 62,000. Occupancy 70%."}), (None, None),
                        industry.measures({"industry_path": ["Hospital"]}))
    assert "INDUSTRY MEASURES (Hospitals): ARPOB (average revenue per occupied bed)" in seen[0]
    ms = out["business"]["measures"]
    assert [m["name"] for m in ms] == ["ARPOB"] and ms[0]["source"]["url"] == "u1"
    assert out["business"]["industry"] == "Hospitals"


def test_deck_shows_measures_and_price_to_book_for_financials():
    v = view_for_deck()
    v["reads"]["business"]["measures"] = [{"name": "ARPOB", "value": "Rs 62,000", "period": "Q1 FY27", "change": None, "quote": "",
                                           "source": {"title": "Deck", "at": "2026-08-01", "url": "u", "kind": "presentation"}}]
    v["reads"]["business"]["industry"] = "Hospitals"
    text = " ".join(sh.text_frame.text for s in Presentation(io.BytesIO(deck.build(v))).slides for sh in s.shapes if sh.has_text_frame)
    assert "Hospitals measures" in text and "P/E" in text
    v["checklist"]["industry"] = {"group": "lender", "label": "Bank or lender", "path": [], "note": ""}
    v["snapshot"]["pb"] = 2.4
    text = " ".join(sh.text_frame.text for s in Presentation(io.BytesIO(deck.build(v))).slides for sh in s.shapes if sh.has_text_frame)
    assert "P/B" in text and "2.4" in text and "P/E" not in text


def test_valuation_fits_the_business():
    p = company()                                       # operating profit isn't in the sample P&L: add it, TTM last
    p["pl"]["rows"]["Operating Profit"] = [15, 18, 22, 27, 31, 38, 40]
    p["balance"]["rows"]["Borrowings"] = [20, 22, 25, 30, 28, 26]
    snap = {"market_cap_cr": 1174, "pe": 45.0, "pb": 3.2}
    v = industry.valuation(p, snap, "general", "hospital")
    assert v["short"] == "EV/EBITDA" and v["value"] == 30.0 and v["pe"] == 45.0      # (1174 + 26) / 40
    assert industry.valuation(p, snap, "lender", "lender")["short"] == "P/B"
    assert industry.valuation(p, snap, "utility", "general")["short"] == "EV/EBITDA"
    assert industry.valuation(p, snap, "general", "it") == {"name": "Price to earnings", "short": "P/E", "value": 45.0, "pe": 45.0,
                                                            "why": "Most businesses are compared on price to earnings."}
    p["pl"]["rows"]["Operating Profit"][-1] = -5        # a loss at the operating level: no multiple, not a negative one
    assert industry.valuation(p, snap, "general", "hospital")["value"] is None


def test_exchange_classification_fills_a_missing_industry():
    import httpx
    from app.intel.filings import NSEFilings

    def handler(r):
        if r.url.path == "/api/quote-equity":
            return httpx.Response(200, json={"industryInfo": {"macro": "Healthcare", "sector": "Healthcare",
                                                              "industry": "Healthcare Services", "basicIndustry": "Hospital"}})
        return httpx.Response(200, text="<html></html>")
    feed = NSEFilings(transport=httpx.MockTransport(handler))
    assert feed.industry("APOLLOHOSP") == ["Healthcare", "Healthcare Services", "Hospital"]
    from app import main
    old = main.filings_feed
    main.filings_feed = feed
    try:
        p = main.with_industry("APOLLOHOSP", {"name": "Apollo"})
        assert industry.measures(p)["key"] == "hospital"
        assert main.with_industry("X", {"industry_path": ["Banks"]})["industry_path"] == ["Banks"]   # the page's own wins
    finally:
        main.filings_feed = old


def test_measures_are_also_read_from_call_transcripts(monkeypatch):
    seen = []

    def complete(system, text, **kw):
        seen.append(text)
        return json.dumps({"summary": "x", "segments": [], "measures": [
            {"name": "Occupancy", "value": "68%", "period": "Q1 FY27", "quote": "Bed occupancy was 68% this quarter", "source": "S2"}]}) \
            if "makes money" in system else json.dumps({"capex": [], "outlook": []})
    monkeypatch.setattr(deepdive, "complete", complete)
    docs = [{"kind": "presentation", "at": "2026-08-01T10:00", "title": "Investor Presentation", "url": "u1"},
            {"kind": "transcript", "at": "2026-08-05T10:00", "title": "Q1 call", "url": "u2"}]
    out = deepdive.read("APOLLOHOSP", "Apollo", "", docs, FakeDocs({"u1": "Hospitals overview.", "u2": "Bed occupancy was 68% this quarter."}),
                        (None, None), industry.measures({"industry_path": ["Hospital"]}))
    assert "[S2] transcript" in seen[0] and "occupancy was 68%" in seen[0]
    assert out["business"]["measures"][0]["source"]["title"] == "Q1 call"



def test_low_promoter_stake_fails_only_when_falling():
    def state(holding):
        p = company()
        p["shareholding"] = {"cols": [f"Q{i}" for i in range(len(holding))], "rows": {"Promoters": holding}}
        return {c["label"]: c for c in checklist.evaluate(p, deepdive.numbers(p))["checks"]}["Promoter holding"]
    assert state([28.2, 28.1, 28.0, 28.0, 28.0])["state"] == "watch"       # Apollo-like: low and steady
    assert "professionally run" in state([28.2, 28.1, 28.0, 28.0, 28.0])["rule"]
    assert state([34, 33, 31, 29, 27.5])["state"] == "fail"                # low and falling
    assert state([25])["state"] == "watch"                                 # no history to judge
    assert state([55, 55, 55, 55, 55])["state"] == "pass"


def test_indian_cash_comes_from_the_other_assets_breakdown():
    import httpx
    from app.intel.screener import Screener, parse
    asked = []

    def handler(r):
        asked.append((r.url.path, dict(r.url.params)))
        if r.url.path == "/api/company/2726/schedules/":
            return httpx.Response(200, json={"Inventories": {"Mar 2024": "100", "Mar 2025": "120"},
                                             "Cash Equivalents": {"Mar 2024": "1,500", "Mar 2025": "2,000"}})
        return httpx.Response(404)
    s = Screener(transport=httpx.MockTransport(handler))
    p = {"company_id": "2726", "basis": "consolidated", "pl": {"cols": ["Mar 2024", "Mar 2025", "TTM"], "rows": {"Operating Profit": [800, 900, 1000], "Net Profit": [400, 450, 500]}},
         "balance": {"cols": ["Mar 2024", "Mar 2025"], "rows": {"Borrowings": [3000, 3500]}}}
    got = s.with_cash(p)
    assert got["balance"]["rows"]["Cash Equivalents"] == [1500, 2000] and asked[0][1]["consolidated"] == "true"
    v = industry.valuation(got, {"market_cap_cr": 10000}, "general", "hospital")
    assert v["short"] == "EV/EBITDA" and v["value"] == round((10000 + 3500 - 2000) / 1000, 1) and "less cash" in v["why"]
    assert s.with_cash(p) == got and len(asked) == 1                      # cached
    # no breakdown (blocked, or not shown to visitors): unchanged, and the note says cash isn't subtracted
    bare = Screener(transport=httpx.MockTransport(lambda r: httpx.Response(403, text="login")))
    assert bare.with_cash(p) is p
    assert "cash isn't subtracted" in industry.valuation(p, {"market_cap_cr": 10000}, "general", "hospital")["why"]
    assert parse('<h1>X</h1><div data-company-id="77" id="company-info"></div>')["company_id"] == "77"
