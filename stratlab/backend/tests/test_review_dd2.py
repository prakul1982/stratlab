"""A harsh audit of the signed-in company pages' screens (DD-001 to DD-072, with DE-016/DE-017): a market mood that only
recited the levels and called an index 2.8% under its high "well below" it, a closed session timed 15:28 on one page and
15:30 on another, alert kinds that don't fit the instrument, a comparison called "slightly ahead" at more than double and
a harsh message for the same company twice, a deck cover from the wrong paragraph with Western grouping and profit as
slivers beside sales, and a watchlist that showed only the day's move."""
import io
import json
from datetime import datetime, timezone

import pytest

from app import deck, main, screens, stock_alerts
from app.intel import ai as A
from app.intel import grounding, routes
from app.intel.news import Wikipedia
from tests import world
from tests.test_investor import view_for_deck


@pytest.fixture
def w(monkeypatch):
    built = world.build(monkeypatch)
    yield built
    built["close"]()


IN = [{"name": "NIFTY 50", "price": 22600.0, "change_pct": -0.39, "high52": 26277.35, "low52": 21743.65},
      {"name": "SENSEX", "price": 73123.83, "change_pct": -0.45, "high52": 85978.25, "low52": 71425.0},
      {"name": "NIFTY BANK", "price": 54287.1, "change_pct": -0.32, "high52": 60000.0, "low52": 50000.0}]
US = [{"name": "S&P 500", "price": 7500.0, "change_pct": -0.5, "high52": 7716.0, "low52": 5000.0},          # 2.8% under its high
      {"name": "NASDAQ", "price": 25000.0, "change_pct": -0.7, "high52": 27500.0, "low52": 17000.0},
      {"name": "DOW JONES", "price": 46000.0, "change_pct": -0.2, "high52": 50000.0, "low52": 38000.0}]
CRUDE = ["Sensex falls as crude oil climbs", "Crude oil at $95 weighs on oil marketing companies", "Oil prices surge on supply fears",
         "TCS Q2 profit up 5%"]


# ---------- DE-016: "well below its 52-week high" only when it is ----------
def test_well_below_a_high_holds_only_for_indices_far_from_it():
    assert not grounding.range_claims_ok("All three indices closed well below their 52-week highs.", US)       # S&P 2.8% under
    assert grounding.range_claims_ok("The Nasdaq closed well below its 52-week high.", US)                     # 9.1% under
    assert not grounding.range_claims_ok("The S&P 500 closed well below its 52-week high.", US)
    assert grounding.range_claims_ok("The Nasdaq stayed far above its 52-week low.", US)
    assert not grounding.range_claims_ok("The NIFTY 50 sits far above its 52-week low.", IN)                    # 3.9% above
    out = grounding.ground_pulse({"tone": "All three indices closed well below their 52-week highs. The S&P 500 fell 0.5%."}, US, [])
    assert "well below" not in out["tone"] and "The S&P 500 fell 0.5%." in out["tone"]


# ---------- DE-017: the index levels in the read as the cards write them ----------
def test_index_levels_in_the_read_are_the_cards_whole_points():
    t = "The NIFTY 50 closed at 22,600.0, down 0.39%, and the SENSEX at 73,123.83; it fell 330.5 points. TCS at ₹3,051.20."
    assert grounding.card_levels(t, IN) == ("The NIFTY 50 closed at 22,600, down 0.39%, and the SENSEX at 73,124; it fell 330.5 points. "
                                            "TCS at ₹3,051.20.")
    assert grounding.card_levels("The S&P 500 ended at 7,499.6.", US) == "The S&P 500 ended at 7,500."
    assert grounding.card_levels("The index gained 22600.4%.", IN) == "The index gained 22600.4%."         # a percentage is never a level


# ---------- DD-001: the mood and the day's story, from the numbers and the headlines ----------
def test_the_mood_is_worked_out_from_the_moves():
    assert grounding.mood_line(IN, {"up": 3, "down": 9, "unchanged": 0, "of": 12}) == \
        "Mildly negative: all three indices down on the day, with 9 of 12 sector indices down."          # the cards carry the moves (MOOD-1)
    mixed = [{**US[0], "change_pct": 0.3}, US[1]]
    assert grounding.mood_line(mixed) == "Mixed: the S&P 500 up and the NASDAQ down on the day."
    assert grounding.mood_line([{**i, "change_pct": 2.4} for i in IN]).startswith("Sharply positive: all three indices up")
    assert grounding.mood_line([{**i, "change_pct": 0.1} for i in IN]).startswith("Flat: all three indices up on the day")
    assert grounding.mood_line([]) == ""


def test_the_main_story_is_the_one_most_headlines_carry():
    assert grounding.main_story("", CRUDE) == "The main story in today's headlines: crude oil prices (3 of 4)."
    # the AI's naming when it is carried as widely; a subject in one headline, or advice, never
    assert grounding.main_story("TCS results", CRUDE) == "The main story in today's headlines: crude oil prices (3 of 4)."
    assert grounding.main_story("buy oil stocks", ["Oil prices rise", "Brent falls"]) == "The main story in today's headlines: crude oil prices (2 of 2)."
    assert grounding.main_story("", ["Gold at a record", "TCS profit up"]) == ""


def test_the_read_says_the_mood_and_where_the_flows_come_from(monkeypatch):
    def fake(system, text, **k):
        assert '"driver"' in system
        return json.dumps({"driver": "crude oil", "tone": "The NIFTY 50 closed at 22,600.0, down 0.39%. FIIs netted an inflow of ₹888 crore.",
                           "hot": [], "flows": [], "themes": []})
    monkeypatch.setattr(A, "complete", fake)
    cash = {"status": "ok", "as_of": "2026-10-23", "fii": {"buy": 9000.0, "sell": 8112.0, "net": 888.0}, "dii": {"buy": 1.0, "sell": 1.0, "net": 3235.0}}
    out = A.pulse("IN", "", IN, [{"headline": h} for h in CRUDE], (None, None), closed=True, cash=cash)
    # the mood is the tables' headline now (mood.py); the read carries the story, and only the headlines' reasons (MOOD-1)
    assert "mood" not in out
    assert out["story_line"] == "The main story in today's headlines: crude oil prices (3 of 4)."
    assert out["story"]["label"] == "crude oil prices" and out["story"]["n"] == 3 and out["story"]["of"] == 4
    assert [out["story"]["headlines"][i] for i in out["story"]["indices"]] == [h for h in CRUDE if "rude" in h or "Brent" in h or "oil" in h.lower()][:3] or len(out["story"]["indices"]) == 3
    assert out["tone"].startswith("The NIFTY 50 closed at 22,600, down")
    assert out["flows_source"] == ("FII and DII figures are the exchange's provisional cash-market numbers for 23 Oct, in ₹ crore; "
                                   "the final ones can differ a little.")
    for vendor in ("NSE", "Yahoo", "Moneycontrol"):
        assert vendor not in out["flows_source"]
    # no FII/DII figure quoted: nothing to source
    monkeypatch.setattr(A, "complete", lambda system, text, **k: json.dumps({"tone": "The NIFTY 50 fell 0.39%.", "hot": [], "flows": [], "themes": []}))
    assert A.pulse("IN", "", IN, [], (None, None), closed=True, cash=cash)["flows_source"] == ""


# ---------- DD-023: a closed session is timed at its close on every page ----------
def test_a_closed_sessions_price_is_as_of_the_close():
    after = datetime(2026, 10, 24, 6, 0, tzinfo=timezone.utc)                  # Saturday
    for at in ("2026-10-23T15:28:41+05:30", "2026-10-23T15:30:00+05:30", "2026-10-23T15:31:00+05:30"):
        assert routes.prices_as_of({"quote": {"price": 1, "at": at}}, "IN", after) == "2026-10-23T15:30+05:30"
    # a stock that last traded at 13:05 says so; no region, the trade's own time (as before)
    assert routes.prices_as_of({"quote": {"price": 1, "at": "2026-10-23T13:05:00+05:30"}}, "IN", after) == "2026-10-23T13:05+05:30"
    assert routes.prices_as_of({"quote": {"price": 1, "at": "2026-10-23T15:28:41+05:30"}}) == "2026-10-23T15:28+05:30"
    # while the session is still on, the trade's time
    during = datetime(2026, 10, 23, 9, 58, tzinfo=timezone.utc)                # 15:28 IST
    assert routes.prices_as_of({"quote": {"price": 1, "at": "2026-10-23T15:28:00+05:30"}}, "IN", during) == "2026-10-23T15:28+05:30"


# ---------- DD-024: alert kinds that fit the instrument ----------
def test_which_alerts_fit_a_symbol(monkeypatch):
    from app import etf_nav, stock_futures
    monkeypatch.setattr(etf_nav, "load_live", lambda: {"rows": {"NIFTYBEES": {"price": 1.0}}})

    class Store:
        def latest(self):
            return ("2026-10-23", {"rows": {"RELIANCE": {}}})
    monkeypatch.setattr(stock_futures.DESK, "store", Store())
    assert stock_alerts.fits("IN", "RELIANCE") == {"etf": False, "fo": True}
    assert stock_alerts.fits("IN", "niftybees") == {"etf": True, "fo": False}
    assert stock_alerts.fits("IN", "POLYCAB") == {"etf": False, "fo": False}
    assert stock_alerts.fits("US", "AAPL") == {"etf": False, "fo": False}
    monkeypatch.setattr(etf_nav, "load_live", lambda: {"rows": {}})            # no list read yet: can't tell
    assert stock_alerts.fits("IN", "RELIANCE")["etf"] is None


def test_the_alert_form_asks_what_fits(w):
    h = world.headers("pro-token")
    r = w["client"].get("/alerts/fits?region=IN&symbol=RELIANCE", headers=h)
    assert r.status_code == 200 and set(r.json()) == {"etf", "fo"}
    assert w["client"].get("/alerts/fits?region=IN&symbol=..%2Fx", headers=h).status_code == 400


# ---------- DD-061: a comparison's "slightly" checked against its figures; the same company twice said kindly ----------
def test_slightly_ahead_needs_a_small_gap():
    s = "Sales and profit growth rates are slightly ahead for TCS at 11.2% and 14.9% vs 10.7% and 6.3% for Infosys."
    assert not grounding.degree_ok(s)
    assert grounding.degree_ok("Sales growth is slightly ahead at 11.2% vs 10.7%.")
    assert grounding.degree_ok("Profit growth is far ahead at 14.9% vs 6.3%.")
    assert not grounding.degree_ok("Profit growth is much higher at 11.2% vs 10.7%.")
    assert grounding.degree_ok("TCS grew slightly faster.")                     # no figures to hold it to


def test_the_same_company_twice_is_a_kind_message(w):
    h = world.headers("pro-token")
    r = w["client"].get("/research/compare?region=IN&a=TCS&b=tcs", headers=h)
    assert r.status_code == 400 and r.json()["detail"]["code"] == "same_symbol"
    assert r.json()["detail"]["message"] == "TCS is in both boxes. Pick a different company in one of them."


# ---------- DD-053: the deck's cover, grouping and charts ----------
def test_the_decks_sales_and_profit_are_two_charts_each_on_its_own_scale():
    slides, _ = deck.plan(view_for_deck())
    s = next(x for x in slides if x.title == "Sales and net profit")
    charts = [b for b in s.blocks if b["kind"] == "columns"]
    assert len(charts) == 2 and all(len(b["series"]) == 1 and b["inr"] for b in charts)
    assert charts[0]["series"][0][0].startswith("Sales (") and charts[1]["series"][0][0].startswith("Net profit (")
    assert charts[0]["box"][0] + charts[0]["box"][2] < charts[1]["box"][0]           # side by side, not overlapping
    assert deck.INDIAN_FORMAT == "[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0"
    assert deck.cap_inr(1610000) == "₹16.1 lakh cr" and deck.cap_inr(4512.3) == "₹4,512 cr"
    from pptx import Presentation
    prs = Presentation(io.BytesIO(deck.build(view_for_deck())))
    assert [sh.has_chart for sh in prs.slides[2].shapes].count(True) == 2
    assert deck.build_pdf(view_for_deck())[:4] == b"%PDF"


def test_the_deck_cover_is_the_company_pages_first_paragraph(monkeypatch):
    v = {**view_for_deck(), "cover": "Acme Industries is India's largest maker of pumps."}
    slides, _ = deck.plan(v)
    assert slides[0].blocks[0]["paras"] == ["Acme Industries is India's largest maker of pumps."]
    # the encyclopedia entry the company page read, its first paragraph; else the profile's first
    wiki = {"extract": "Acme Industries is an Indian maker of pumps, founded in 1950.\nIt was founded by A. Acme in Pune."}
    monkeypatch.setattr(main.research_hub.wiki, "company", lambda name, cached=False: wiki if cached else pytest.fail("fetched"))
    assert main.deck_cover("ACME", {"p": {"name": "Acme Industries", "about": "Acme was founded by A. Acme."}}) == \
        "Acme Industries is an Indian maker of pumps, founded in 1950."
    monkeypatch.setattr(main.research_hub.wiki, "company", lambda name, cached=False: None)
    assert main.deck_cover("ACME", {"p": {"name": "Acme", "about": "Acme was founded by A. Acme in Pune in 1950."}}) == \
        "Acme was founded by A. Acme in Pune in 1950."


def test_a_cached_only_encyclopedia_lookup_never_calls_the_site():
    import httpx

    def boom(request):
        raise AssertionError("no call")
    wk = Wikipedia(transport=httpx.MockTransport(boom))
    assert wk.company("Acme Industries", cached=True) is None
    wk.cache.set(("/api/rest_v1/page/summary/Acme_Industries", (), "json"), {"title": "Acme Industries", "extract": "Pumps."}, 600)
    wk.cache.set(("/w/api.php", tuple(sorted({"action": "query", "list": "search", "srsearch": "Acme Industries company", "srlimit": 1,
                                                "format": "json"}.items())), "json"), {"query": {"search": [{"title": "Acme Industries"}]}}, 600)
    assert wk.company("Acme Industries", cached=True)["extract"] == "Pumps."


# ---------- DD-072: the watchlist's list view has the P/E and the 52-week position ----------
def test_the_watchlist_facts_come_from_the_screeners_index(w, monkeypatch):
    rows = [{"symbol": "TCS", "pe": 22.41, "from_high": -18.3, "price_at": "2026-10-23"}, {"symbol": "ITC", "pe": None, "from_high": -0.2}]
    monkeypatch.setattr(screens, "load_index", lambda region: {"region": region, "rows": rows, "at": "2026-10-23T12:00:00+00:00"})
    h = world.headers("pro-token")
    got = w["client"].get("/research/watchlist/facts?region=IN&symbols=TCS,itc,NOPE", headers=h).json()
    assert got["facts"] == {"TCS": {"pe": 22.41, "from_high": -18.3, "price_at": "2026-10-23"},
                            "ITC": {"pe": None, "from_high": -0.2, "price_at": None}}
    assert got["as_of"] == "2026-10-23"
